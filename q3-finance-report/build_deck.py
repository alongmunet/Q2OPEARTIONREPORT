import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
sys.path.insert(0, __import__("os").path.dirname(__file__))
from q3data import *

OUT = sys.argv[1]
FONT = "Aptos"

NAVY, GOLD, ORANGE, GREEN = "081E42", "FFC20E", "F79009", "22C55E"
BG, LGOLD, ALT, WHITE, PCT = "F6F8FB", "FFF2B3", "F8FAFC", "FFFFFF", "DDF0F8"
RED, BORDER, INK, MUTED = "C00000", "E2E8F0", "0F172A", "64748B"

def rgb(h): return RGBColor.from_string(h)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
L, W = 0.28, 13.333 - 0.56

def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(0.75)
    else: s.line.fill.background()
    st = s._element.find(qn('p:style'))
    if st is not None: s._element.remove(st)
    return s

def text(slide, x, y, w, h, runs, size=10, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    """runs: str, or list of paragraphs; each paragraph a list of (text, bold, color[, size])."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.02); tf.margin_top = tf.margin_bottom = Inches(0.01)
    if isinstance(runs, str): runs = [[(runs, False, INK)]]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for r in para:
            t, b, c = r[:3]
            run = p.add_run(); run.text = t
            f = run.font; f.name = FONT; f.bold = b; f.color.rgb = rgb(c); f.size = Pt(r[3] if len(r) > 3 else size)
    return tb

def frame(title, headline, headline_size=11):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    rect(s, 0, 0, 13.333, 0.84, NAVY)
    rect(s, 0, 0.84, 13.333, 0.11, GOLD)
    text(s, L, 0.12, 10, 0.45, [[(title, True, WHITE, 24)]])
    text(s, L, 0.56, 10, 0.2, [[("TGFB | July 2026 - September 2026 | All figures in RM", False, WHITE, 7)]])
    badge = rect(s, 10.89, 0.2, 1.33, 0.43, GOLD)
    tf = badge.text_frame; tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "Q3 2026"; r.font.name = FONT; r.font.bold = True; r.font.size = Pt(9); r.font.color.rgb = rgb(NAVY)
    text(s, L, 1.13, W, 0.32, [[(headline, True, INK, headline_size)]])
    text(s, L, 7.3, 8, 0.15, [[("Source: TGFB Cash Flow Statement PDFs for July, August and September 2026", False, MUTED, 5.5)]])
    return s

def section_bar(s, y, label, h=0.24, size=10.5):
    bar = rect(s, L, y, W, h, NAVY)
    text(s, L + 0.04, y, W, h, [[(label, True, WHITE, size)]], anchor=MSO_ANCHOR.MIDDLE)

def money(v, neg_paren=True):
    if abs(v) < 0.005: return "0.00"
    return f"({abs(v):,.2f})" if (v < 0 and neg_paren) else f"{v:,.2f}"

def M(v):   # RM millions for headlines
    return f"RM{abs(v)/1e6:.2f}M"
def K(v):
    return f"RM{abs(v)/1e3:.1f}K"

def set_border(cell, color=WHITE, width=Pt(1)):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None: tcPr.remove(old)
        ln = etree.SubElement(tcPr, qn(tag), w=str(int(width)), cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val=color)
        etree.SubElement(ln, qn("a:prstDash"), val="solid")
    # schema order: borders must precede fill
    fill = tcPr.find(qn("a:solidFill"))
    if fill is not None: tcPr.remove(fill); tcPr.append(fill)

def table(s, x, y, col_w, row_h, header, rows, size=9, header_size=None, first_bold=True):
    """rows: list of lists of (text, fill, color, bold, align)."""
    n = len(rows) + 1
    gt = s.shapes.add_table(n, len(col_w), Inches(x), Inches(y), Inches(sum(col_w)), Inches(row_h[0] + sum(row_h[1:])))
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    for a in ("firstRow", "bandRow"): tblPr.set(a, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None: sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for i, w in enumerate(col_w): tbl.columns[i].width = Inches(w)
    for i, h in enumerate(row_h): tbl.rows[i].height = Inches(h)
    all_rows = [[(h, GOLD, NAVY, True, PP_ALIGN.CENTER) for h in header]] + rows
    for ri, row in enumerate(all_rows):
        for ci, (t, fill, color, bold, align) in enumerate(row):
            c = tbl.cell(ri, ci)
            c.fill.solid(); c.fill.fore_color.rgb = rgb(fill)
            c.margin_left = c.margin_right = Inches(0.05); c.margin_top = c.margin_bottom = Inches(0.0)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.alignment = align
            p.font.size = Pt(header_size if (ri == 0 and header_size) else size)
            r = p.add_run(); r.text = t or " "
            r.font.name = FONT; r.font.bold = bold; r.font.color.rgb = rgb(color)
            r.font.size = Pt(header_size if (ri == 0 and header_size) else size)
            set_border(c)
    return tbl

R, C, Lf = PP_ALIGN.RIGHT, PP_ALIGN.CENTER, PP_ALIGN.LEFT
MONTH_FILLS = [ALT, WHITE, ALT]

def num_cells(vals, total, fills=MONTH_FILLS, total_fill=WHITE, neg_red=True, paren=True, base_color=INK):
    out = []
    for v, f in zip(list(vals) + [total], list(fills) + [total_fill]):
        out.append((money(v, paren), f, RED if (neg_red and v < -0.005) else base_color, True, R))
    return out

def data_row(label, vals, pct, remark, signed=True):
    return ([(label, LGOLD, INK, True, Lf)] + num_cells(vals, sum(vals), neg_red=signed)
            + [(f"{pct:.1%}", PCT, INK, True, C), (remark, WHITE, INK, True, Lf)])

def total_row(label, vals, pct, remark):
    cells = [(label, NAVY, WHITE, True, Lf)]
    for v in list(vals) + [sum(vals)]:
        cells.append((money(v), GOLD, RED if v < -0.005 else INK, True, R))
    if pct is not None: cells.append((f"{pct:.1%}", GOLD, INK, True, C))
    cells.append((remark, GOLD, INK, True, Lf))
    return cells

# ---------- derived figures ----------
paid_m = [sum(v[i] for _, v in PAID) for i in range(3)]
rec_q, paid_q = sum(RECEIPTS_TOTAL), sum(paid_m)
ops_q, inv_q, fin_q, net_q = sum(NET_OPS), sum(NET_INV), sum(NET_FIN), sum(NET_CASH)
rd = dict(RECEIPTS); pdd = dict(PAID); invd = dict(INVESTING)
cust_share = sum(rd["Customers & Sales"]) / rec_q
capex = -sum(invd["Capex New Branches"])
sst = sum(pdd["SST"])

# ---------- Slide 1: Cash Flow Comparison ----------
s = frame("Financial Performance: Cash Flow Comparison",
          f"Q3 operating cash was positive overall at {M(ops_q)}, but negative in July after a {M(pdd['SST'][0])} SST payment. "
          f"Net cash flow was ({M(net_q)}), mainly from {M(capex)} new-branch capex and a {M(2e6)} fixed deposit placement.", 10.5)
cards = [
    ("Q3 CASH RECEIPTS", M(rec_q), INK, f"Customers & sales contribute {cust_share:.1%}", GOLD),
    ("Q3 CASH PAID", M(paid_q), INK, "Inventory + wages are key outflows", ORANGE),
    ("NET CASH FROM OPERATIONS", M(ops_q), INK, "Negative in Jul; positive Aug-Sep", GREEN),
    ("Q3 NET CASH FLOW", f"({M(net_q)})", RED, "Quarter total movement", NAVY),
    ("SEPT NET CASH FLOW", f"({M(NET_CASH[2])})", RED, "Sep-26 monthly movement", GOLD),
]
cw, gap = 2.27, 0.355
for i, (lab, val, vc, sub, bar) in enumerate(cards):
    x = L + i * (cw + gap)
    rect(s, x, 1.6, cw, 0.78, WHITE, BORDER)
    rect(s, x, 1.6, 0.06, 0.78, bar)
    text(s, x + 0.14, 1.66, cw - 0.2, 0.16, [[(lab, False, MUTED, 6.5)]])
    text(s, x + 0.14, 1.8, cw - 0.2, 0.34, [[(val, True, vc, 17)]])
    text(s, x + 0.14, 2.15, cw - 0.2, 0.16, [[(sub, False, MUTED, 6)]])

section_bar(s, 2.56, "Monthly Cash Flow Summary", h=0.26, size=10)
cols1 = [2.2, 1.45, 1.45, 1.45, 1.55, 4.67]
def s1row(label, vals, remark, paid=False):
    v = [-x for x in vals] if paid else vals
    return [(label, LGOLD, INK, True, Lf)] + num_cells(v, sum(v)) + [(remark, WHITE, INK, True, C)]
rows1 = [
    s1row("Cash Receipts", RECEIPTS_TOTAL, "Increased every month; Sep highest at " + M(RECEIPTS_TOTAL[2])),
    s1row("Cash Paid", paid_m, f"Jul highest ({M(pdd['SST'][0])} SST); Aug lowest with no SST payment"),
    s1row("Net Cash from Operations", NET_OPS, f"Negative in Jul; Aug strongest at {M(NET_OPS[1])}"),
    s1row("Net Investing Activities", NET_INV, f"Jul capex {M(invd['Capex New Branches'][0])}; Sep includes {M(2e6)} fixed deposit placement"),
    s1row("Net Financing Activities", NET_FIN, f"Aug positive from {K(230281.85)} repayment from related parties"),
]
net1 = [("Net Cash Flow", NAVY, WHITE, True, Lf)] + [(money(v), GOLD, RED if v < 0 else INK, True, R) for v in NET_CASH + [net_q]] \
       + [(f"Bank balance {M(OPENING[0])} (1 Jul) to {M(CLOSING[2])} (30 Sep)", WHITE, INK, True, C)]
table(s, L, 2.82, cols1, [0.4] + [0.47] * 6, ["Category", "Jul-26", "Aug-26", "Sep-26", "Q3 Total", "Remarks"],
      rows1 + [net1], size=10, header_size=10)
rect(s, L, 6.35, W, 0.48, WHITE, BORDER)
text(s, L + 0.12, 6.35, W - 0.24, 0.48,
     [[("KEY TAKEAWAY ", True, MUTED, 8),
       (f"Cash receipts grew every month from {M(RECEIPTS_TOTAL[0])} to {M(RECEIPTS_TOTAL[2])}, but SST payments in July and September ({M(sst)} combined) "
        f"and new-branch capex of {M(capex)} pulled cash down. {M(2e6)} of the decline was placed in fixed deposit, not spent.", True, INK, 8)]],
     anchor=MSO_ANCHOR.MIDDLE)

# ---------- Slide 2: Cash Receipts ----------
s = frame("Financial Performance: Cash Receipts Comparison",
          f"Total cash receipts increased every month, from {M(RECEIPTS_TOTAL[0])} in July to {M(RECEIPTS_TOTAL[2])} in September. "
          "Royalty declined after July.", 12)
section_bar(s, 1.68, "Operations: Cash Receipts From", h=0.36, size=11)
cols2 = [2.6, 1.3, 1.3, 1.3, 1.5, 0.85, 3.92]
rem2 = {
    "Customers & Sales": "Grew every month; Sep-26 highest",
    "Licensing fee & package": "No receipts in Jul; RM120K licensing fee in Sep",
    "Royalty": "Declined Aug-Sep, in line with royalty discount",
    "Other operations": "Stable at about RM42K every month",
}
rows2 = [data_row(n, v, sum(v) / rec_q, rem2[n]) for n, v in RECEIPTS]
rows2.append(total_row("TOTAL CASH RECEIPTS", RECEIPTS_TOTAL, 1.0, "Sep highest; Jul lowest"))
table(s, L, 2.04, cols2, [0.5] + [0.64] * 5, ["Account", "Jul-26", "Aug-26", "Sep-26", "Q3 Total", "% of Q3", "Compare / Remarks"],
      rows2, size=10.5, header_size=10.5)
rect(s, L, 6.0, 6.5, 1.12, WHITE, BORDER)
roy = rd["Royalty"]; lic = rd["Licensing fee & package"]; oth = rd["Other operations"]
text(s, L + 0.2, 6.1, 6.1, 0.95, [
    [("Receipts insights", True, INK, 10)],
    [(f"• Customers & Sales contributed {cust_share:.1%} of Q3 receipts.", False, INK, 8.5)],
    [(f"• Royalty fell from {K(roy[0])} in July to {K(roy[2])} in September.", False, INK, 8.5)],
    [(f"• Licensing fee receipts rose from nil in July to {K(lic[2])} in September.", False, INK, 8.5)],
    [(f"• Other operations stayed stable at {K(min(oth))}-{K(max(oth))} per month.", False, INK, 8.5)],
])

# ---------- Slide 3: Operating Cash Paid ----------
s = frame("Financial Performance: Operating Cash Paid Comparison",
          "Inventory purchases, wages, SST and rental were the biggest operating cash outflows. "
          "SST was paid in July and September, making them the two highest-outflow months.", 10)
section_bar(s, 1.5, "Operations: Cash Paid For", h=0.24, size=11)
cols3 = [3.0, 1.35, 1.35, 1.35, 1.45, 0.75, 3.52]
rem3 = {
    "Inventory purchases": "Peaked Aug; largest outflow",
    "Distribution Fee": "Sep payment only",
    "Deposit": "Sep-26 highest; Aug-26 lowest",
    "HQ Operation": "Increased every month",
    "Operation": "No payment in Q3",
    "Marketing": "Aug-26 highest; Sep-26 lowest",
    "Profit sharing": "Aug-26 highest; Jul-26 lowest",
    "Wage expenses": "Stable at RM1.78M-RM1.83M monthly",
    "Bank Charges": "Minimal",
    "Income taxes": "Stable monthly RM250K",
    "New Branches Expenses": "Reduced every month",
    "Branch Expenditure & Upkeep": "Increased gradually every month",
    "Credit Card": "Sep-26 highest; Aug-26 lowest",
    "Open House & Iftar": "No payment in Q3",
    "Co Trip": "Aug payment only",
    "Utilities": "Aug-26 highest; Jul-26 lowest",
    "Rental": "Increased every month",
    "Sponsorship": "Jul payment only",
    "SST": "Paid in Jul and Sep; none in Aug",
    "Donation": "No payment in Q3",
}
rows3 = [data_row(n, v, sum(v) / paid_q, rem3[n], signed=False) for n, v in PAID]
rows3.append(total_row("TOTAL CASH PAID", paid_m, 1.0, "Jul highest; Aug lowest (no SST)"))
# totals row shows outflows as positive, like the template
rows3[-1] = [rows3[-1][0]] + [(money(v), GOLD, INK, True, R) for v in paid_m + [paid_q]] + rows3[-1][5:]
table(s, L, 1.74, cols3, [0.22] + [0.205] * 21, ["Account", "Jul-26", "Aug-26", "Sep-26", "Q3 Total", "% Paid", "Compare / Remarks"],
      rows3, size=8.5, header_size=8.5)
top = sorted(PAID, key=lambda kv: -sum(kv[1]))[:5]
rect(s, L, 6.6, W, 0.48, WHITE, BORDER)
text(s, L + 0.12, 6.6, W - 0.24, 0.48,
     [[("KEY TAKEAWAY ", True, MUTED, 8),
       ("TOP OUTFLOW CATEGORIES: " + " | ".join(f"{n} {M(sum(v))}" for n, v in top), True, INK, 8)]],
     anchor=MSO_ANCHOR.MIDDLE)

# ---------- Slide 4: Investing & Financing ----------
s = frame("Financial Performance: Investing & Financing Comparison",
          f"New-branch capex ({M(capex)}) and a {M(2e6)} fixed deposit placement in September drove investing outflows. "
          "Financing was positive only in August, from related-party repayments.", 9.5)
cols4 = [3.0, 1.35, 1.35, 1.35, 1.45, 0.75, 3.52]
hdr4 = ["Account", "Jul-26", "Aug-26", "Sep-26", "Q3 Total", "% / Q3", "Compare / Remarks"]
inv_abs = sum(abs(x) for _, v in INVESTING for x in v)
rem_inv = {
    "Proceeds from sale of motor vehicle": "Disposal of motor vehicle in Jul",
    "Interest received (Maybank pooled)": "Interest from Maybank pooled account",
    "Collection of Investor funds": "Investor fund; Sep received via RHB",
    "Sale of investment securities": "",
    "Purchase of property and equipment": "Reduced every month",
    "Capex New Branches": "For outlet expansion; Jul-26 highest",
    "Placement of Fixed Deposit": "RM2.0M placed in fixed deposit in Sep",
    "Purchase of investment": "",
}
section_bar(s, 1.48, "Investing Activities", h=0.22, size=10.5)
rows4 = [data_row(n, v, sum(abs(x) for x in v) / inv_abs, rem_inv[n]) for n, v in INVESTING]
rows4.append(total_row("NET INVESTING ACTIVITIES", NET_INV, 1.0, "Sep-26 highest outflow; Aug-26 lowest"))
table(s, L, 1.7, cols4, [0.2] * 10, hdr4, rows4, size=8, header_size=8)

fin_abs = sum(abs(x) for _, v in FINANCING for x in v)
rem_fin = {
    "Issuance of stock": "", "Loan received": "No new loan in Q3",
    "Collection of Investor funds": "Investor fund recorded under financing in Sep",
    "Borrowing": "",
    "Repayment from related parties": f"Net outflow in Jul; {K(230281.85)} received in Aug",
    "Purchase of investment": "Jul only (recorded under financing)",
    "Repayment of loans": "Stable at about RM155K-RM160K monthly",
    "Dividends": "",
    "Advance to related parties": "Jul-26 highest; Aug-26 minimal",
    "Repayment to related parties": "", "Pending": "",
}
section_bar(s, 3.86, "Financing Activities", h=0.22, size=10.5)
rows5 = [data_row(n, v, sum(abs(x) for x in v) / fin_abs, rem_fin[n]) for n, v in FINANCING]
rows5.append(total_row("NET FINANCING ACTIVITIES", NET_FIN, 1.0, "Aug-26 positive; Jul-26 lowest"))
table(s, L, 4.08, cols4, [0.2] * 13, hdr4, rows5, size=8, header_size=8)

prs.save(OUT)
print("saved", OUT, "| rec", rec_q, "paid", paid_q, "capex", capex, "sst", sst, "top", [(n, round(sum(v))) for n, v in top])
