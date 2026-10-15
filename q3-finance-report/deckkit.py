"""Shared slide helpers for the TGFB finance decks (navy/gold Q2 finance report template)."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
FONT = "Aptos"

# Per-deck header/footer text; set these before calling frame().
SUBTITLE = ""
BADGE = ""
SOURCE = ""

NAVY, GOLD, ORANGE, GREEN = "081E42", "FFC20E", "F79009", "22C55E"
BG, LGOLD, ALT, WHITE, PCT = "F6F8FB", "FFF2B3", "F8FAFC", "FFFFFF", "DDF0F8"
RED, BORDER, INK, MUTED = "C00000", "E2E8F0", "0F172A", "64748B"

def rgb(h): return RGBColor.from_string(h)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
L, W = 0.28, 13.333 - 0.56

def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(0.75)
    else: s.line.fill.background()
    st = s._element.find(qn('p:style'))
    if st is not None: s._element.remove(st)
    return s

def text(slide, x, y, w, h, runs, size=10, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    """runs: str, or list of paragraphs; each paragraph a list of (text, bold, color[, size])."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.02); tf.margin_top = tf.margin_bottom = Inches(0.01)
    if isinstance(runs, str): runs = [[(runs, False, INK)]]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for r in para:
            t, b, c = r[:3]
            run = p.add_run(); run.text = t
            f = run.font; f.name = FONT; f.bold = b; f.color.rgb = rgb(c); f.size = Pt(r[3] if len(r) > 3 else size)
    return tb

def frame(title, headline, headline_size=11):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    rect(s, 0, 0, 13.333, 0.84, NAVY)
    rect(s, 0, 0.84, 13.333, 0.11, GOLD)
    text(s, L, 0.12, 10, 0.45, [[(title, True, WHITE, 24)]])
    text(s, L, 0.56, 10, 0.2, [[(SUBTITLE, False, WHITE, 7)]])
    badge = rect(s, 10.89, 0.2, 1.33, 0.43, GOLD)
    tf = badge.text_frame; tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = BADGE; r.font.name = FONT; r.font.bold = True; r.font.size = Pt(9); r.font.color.rgb = rgb(NAVY)
    text(s, L, 1.13, W, 0.32, [[(headline, True, INK, headline_size)]])
    text(s, L, 7.3, 12, 0.15, [[(SOURCE, False, MUTED, 5.5)]])
    return s

def section_bar(s, y, label, h=0.24, size=10.5):
    bar = rect(s, L, y, W, h, NAVY)
    text(s, L + 0.04, y, W, h, [[(label, True, WHITE, size)]], anchor=MSO_ANCHOR.MIDDLE)

def money(v, neg_paren=True):
    if abs(v) < 0.005: return "0.00"
    return f"({abs(v):,.2f})" if (v < 0 and neg_paren) else f"{v:,.2f}"

def M(v):   # RM millions for headlines
    return f"RM{abs(v)/1e6:.2f}M"
def K(v):
    return f"RM{abs(v)/1e3:.1f}K"

def set_border(cell, color=WHITE, width=Pt(1)):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None: tcPr.remove(old)
        ln = etree.SubElement(tcPr, qn(tag), w=str(int(width)), cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val=color)
        etree.SubElement(ln, qn("a:prstDash"), val="solid")
    # schema order: borders must precede fill
    fill = tcPr.find(qn("a:solidFill"))
    if fill is not None: tcPr.remove(fill); tcPr.append(fill)

def table(s, x, y, col_w, row_h, header, rows, size=9, header_size=None, first_bold=True):
    """rows: list of lists of (text, fill, color, bold, align)."""
    n = len(rows) + 1
    gt = s.shapes.add_table(n, len(col_w), Inches(x), Inches(y), Inches(sum(col_w)), Inches(row_h[0] + sum(row_h[1:])))
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    for a in ("firstRow", "bandRow"): tblPr.set(a, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None: sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for i, w in enumerate(col_w): tbl.columns[i].width = Inches(w)
    for i, h in enumerate(row_h): tbl.rows[i].height = Inches(h)
    all_rows = [[(h, GOLD, NAVY, True, PP_ALIGN.CENTER) for h in header]] + rows
    for ri, row in enumerate(all_rows):
        for ci, (t, fill, color, bold, align) in enumerate(row):
            c = tbl.cell(ri, ci)
            c.fill.solid(); c.fill.fore_color.rgb = rgb(fill)
            c.margin_left = c.margin_right = Inches(0.05); c.margin_top = c.margin_bottom = Inches(0.0)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.alignment = align
            p.font.size = Pt(header_size if (ri == 0 and header_size) else size)
            r = p.add_run(); r.text = t or " "
            r.font.name = FONT; r.font.bold = bold; r.font.color.rgb = rgb(color)
            r.font.size = Pt(header_size if (ri == 0 and header_size) else size)
            set_border(c)
    return tbl

R, C, Lf = PP_ALIGN.RIGHT, PP_ALIGN.CENTER, PP_ALIGN.LEFT
MONTH_FILLS = [ALT, WHITE, ALT]

def num_cells(vals, total, fills=MONTH_FILLS, total_fill=WHITE, neg_red=True, paren=True, base_color=INK):
    out = []
    for v, f in zip(list(vals) + [total], list(fills) + [total_fill]):
        out.append((money(v, paren), f, RED if (neg_red and v < -0.005) else base_color, True, R))
    return out

def data_row(label, vals, pct, remark, signed=True):
    return ([(label, LGOLD, INK, True, Lf)] + num_cells(vals, sum(vals), neg_red=signed)
            + [(f"{pct:.1%}", PCT, INK, True, C), (remark, WHITE, INK, True, Lf)])

def total_row(label, vals, pct, remark):
    cells = [(label, NAVY, WHITE, True, Lf)]
    for v in list(vals) + [sum(vals)]:
        cells.append((money(v), GOLD, RED if v < -0.005 else INK, True, R))
    if pct is not None: cells.append((f"{pct:.1%}", GOLD, INK, True, C))
    cells.append((remark, GOLD, INK, True, Lf))
    return cells

